import json
import logging
import os
import uuid
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Form
from slowapi import Limiter
from slowapi.util import get_remote_address
from fastapi.responses import Response
from pydantic import BaseModel
import google.generativeai as genai
from google.generativeai import protos as genai_protos
from google.protobuf.struct_pb2 import Struct
from sqlalchemy.orm import Session

from app.config import GEMINI_API_KEY, MAX_UPLOAD_SIZE_BYTES
from app.auth import get_current_user
from app.models.user import User
from app.models.chat import ChatSession, ChatMessage
from app.database import get_db
from app.services.ai_tools import AVAILABLE_TOOLS, TOOL_FUNCTIONS, set_ai_user_context, clear_ai_user_context

router = APIRouter(prefix="/api/ai", tags=["Assistente IA"])
_ai_limiter = Limiter(key_func=get_remote_address)

# Diretório para uploads e documentos gerados
UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

# Endpoints de debug removidos por segurança (test_pdf, install_deps)

# Inicializa as credenciais da Gemini
if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)

class ChatMessageDTO(BaseModel):
    role: str
    content: str
    
class ChatRequest(BaseModel):
    session_id: Optional[int] = None
    message: str
    file_context: Optional[str] = None  # Texto extraído do arquivo enviado

class SessionResponse(BaseModel):
    id: int
    titulo: str
    created_at: str
    updated_at: str
    message_count: int

SYSTEM_PROMPT = """
Você é a **Pepétua**, a inteligência artificial do sistema "CRM Brasileiros no Atacama", promovida a modo DEUS (God Mode).
Seu papel é ajudar o usuário administrando leads, viagens, tarefas, equipes, etc.

COMO COMPLETAR SUAS TAREFAS:
Você tem duas abordagens principais para alterar ou consultar o sistema:
1. **Agentic API (PREFERENCIAL)**: Você já tem mapeado em sua memória (veja abaixo) todas as rotas /api/ disponíveis. Use a ferramenta `call_internal_api` para simular requisições HTTP para essas rotas. Use preferencialmente as rotas oficias (ex: `POST /api/leads`) porque elas aplicam validações de campo, gatilham automações N8N e garantem a integridade dos dados!
2. **God Mode SQL**: Se a rota não existir, não funcionar, ou for uma alteração em massa complexa, você tem as ferramentas `run_select_query` e `run_sql_write_query` para modificar **qualquer** tabela diretamente no banco de dados (`leads`, `tags`, etc). Use `get_database_schema` se não souber os campos. NUNCA modifique a tabela `users` por SQL.
3. Outros helpers fáceis: `create_lead`, `add_tag_to_lead`.

CAPACIDADES DE DOCUMENTOS:
- Você pode **ler** arquivos enviados pelo usuário (PDF, Excel, CSV, DOCX, TXT). O texto extraído será enviado junto com a mensagem.
- Você pode **gerar** documentos para o usuário. Use as ferramentas `generate_excel_document` e `generate_pdf_document` para criar relatórios, listas e outros documentos sob demanda.
- REGRA CRÍTICA PARA LINKS DE DOWNLOAD: Quando as ferramentas de geração retornarem o campo "download_url", você DEVE usar EXATAMENTE esse valor como link. NÃO invente URLs, NÃO adicione domínio, NÃO mude o caminho. Use sempre o link relativo EXATO retornado pela ferramenta.
- Formato OBRIGATÓRIO do link de download no markdown: [📥 Baixar arquivo](/api/ai/download/NOME_DO_ARQUIVO.extensao)
- Exemplo correto: [📥 Baixar arquivo](/api/ai/download/relatorio_leads_abc123.xlsx)
- ❌ NUNCA faça isso: [Baixar](https://exemplo.com/download/arquivo) — isso vai QUEBRAR o download!
"""

# ─── Session Management ──────────────────────────────────────────────

@router.get("/sessions", summary="Listar sessões de chat do usuário")
def list_sessions(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Retorna todas as sessões de chat do usuário logado, ordenadas pela mais recente."""
    sessions = (
        db.query(ChatSession)
        .filter(ChatSession.user_id == current_user.id)
        .order_by(ChatSession.updated_at.desc())
        .all()
    )
    return {
        "sessions": [
            {
                "id": s.id,
                "titulo": s.titulo,
                "created_at": str(s.created_at),
                "updated_at": str(s.updated_at),
                "message_count": len(s.messages),
            }
            for s in sessions
        ]
    }

@router.post("/sessions", summary="Criar nova sessão de chat")
def create_session(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Cria uma nova sessão de chat vazia."""
    session = ChatSession(user_id=current_user.id, titulo="Nova conversa")
    db.add(session)
    db.commit()
    db.refresh(session)
    return {"id": session.id, "titulo": session.titulo}

@router.get("/sessions/{session_id}/messages", summary="Buscar mensagens de uma sessão")
def get_session_messages(
    session_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Retorna todas as mensagens de uma sessão em ordem cronológica."""
    session = (
        db.query(ChatSession)
        .filter(ChatSession.id == session_id, ChatSession.user_id == current_user.id)
        .first()
    )
    if not session:
        raise HTTPException(status_code=404, detail="Sessão não encontrada")
    
    return {
        "session_id": session.id,
        "titulo": session.titulo,
        "messages": [
            {"role": m.role, "content": m.content, "created_at": str(m.created_at)}
            for m in session.messages
        ]
    }

@router.delete("/sessions/{session_id}", summary="Excluir sessão de chat")
def delete_session(
    session_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Exclui uma sessão de chat e todas as suas mensagens."""
    session = (
        db.query(ChatSession)
        .filter(ChatSession.id == session_id, ChatSession.user_id == current_user.id)
        .first()
    )
    if not session:
        raise HTTPException(status_code=404, detail="Sessão não encontrada")
    
    db.delete(session)
    db.commit()
    return {"message": "Sessão excluída"}

# ─── File Upload ─────────────────────────────────────────────────────

@router.post("/upload", summary="Upload de arquivo para a IA ler")
async def upload_file_for_ai(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    """
    Faz upload de um arquivo e extrai o texto para a IA ler.
    Formatos suportados: PDF, Excel (.xlsx/.xls), CSV, DOCX, TXT
    """
    filename = file.filename.lower() if file.filename else ""
    content = await file.read()
    
    # Limitar tamanho do upload
    if len(content) > MAX_UPLOAD_SIZE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"Arquivo muito grande. Máximo permitido: {MAX_UPLOAD_SIZE_BYTES // (1024*1024)}MB"
        )
    
    extracted_text = ""
    
    try:
        if filename.endswith(".pdf"):
            import pdfplumber
            import io
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                pages_text = []
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text:
                        pages_text.append(f"--- Página {i+1} ---\n{text}")
                extracted_text = "\n\n".join(pages_text)
        
        elif filename.endswith((".xlsx", ".xls")):
            import openpyxl
            import io
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets_text = []
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    sheet_lines = [f"--- Planilha: {ws.title} ---"]
                    for row in rows:
                        line = " | ".join(str(cell) if cell is not None else "" for cell in row)
                        sheet_lines.append(line)
                    sheets_text.append("\n".join(sheet_lines))
            wb.close()
            extracted_text = "\n\n".join(sheets_text)
        
        elif filename.endswith(".csv"):
            import csv
            import io
            text = content.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text))
            lines = []
            for row in reader:
                lines.append(" | ".join(row))
            extracted_text = "\n".join(lines)
        
        elif filename.endswith(".docx"):
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            extracted_text = "\n".join(paragraphs)
        
        elif filename.endswith(".txt"):
            try:
                extracted_text = content.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = content.decode("latin-1")
        
        else:
            raise HTTPException(
                status_code=400,
                detail="Formato não suportado. Use PDF, Excel, CSV, DOCX ou TXT."
            )
        
        if not extracted_text.strip():
            extracted_text = "(Arquivo vazio ou sem texto extraível)"
        
        # Limitar tamanho do texto para não estourar o contexto da IA
        MAX_CHARS = 50000
        if len(extracted_text) > MAX_CHARS:
            extracted_text = extracted_text[:MAX_CHARS] + f"\n\n... (texto truncado, {len(extracted_text)} caracteres no total)"
        
        return {
            "filename": file.filename,
            "chars_extracted": len(extracted_text),
            "text": extracted_text
        }
    
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao processar arquivo: {str(e)}")

# ─── File Download ───────────────────────────────────────────────────

@router.get("/download/{filename}", summary="Baixar arquivo gerado pela IA",
            response_class=Response)
async def download_file(
    filename: str,
    current_user: User = Depends(get_current_user),
):
    """Baixa um arquivo gerado pela IA com validações robustas. Requer autenticação."""
    import logging
    logger = logging.getLogger("ai.download")
    
    # 1. Sanitizar filename — prevenir path traversal
    safe_name = os.path.basename(filename)
    if safe_name != filename or ".." in filename:
        logger.warning(f"[DOWNLOAD] Tentativa de path traversal bloqueada: {filename}")
        raise HTTPException(status_code=400, detail="Nome de arquivo inválido")
    
    filepath = os.path.join(UPLOAD_DIR, safe_name)
    
    # 2. Verificar existência
    if not os.path.isfile(filepath):
        logger.error(f"[DOWNLOAD] Arquivo não encontrado: {filepath}")
        try:
            logger.error(f"[DOWNLOAD] Conteúdo de uploads/: {os.listdir(UPLOAD_DIR)}")
        except Exception:
            pass
        raise HTTPException(status_code=404, detail=f"Arquivo '{safe_name}' não encontrado no servidor")
    
    # 3. Verificar tamanho mínimo
    file_size = os.path.getsize(filepath)
    if file_size == 0:
        logger.error(f"[DOWNLOAD] Arquivo vazio: {filepath}")
        raise HTTPException(status_code=500, detail="Arquivo está vazio no servidor")
    
    logger.info(f"[DOWNLOAD] Servindo arquivo: {safe_name} ({file_size} bytes)")
    print(f"[DOWNLOAD] Servindo: {safe_name} | {file_size} bytes | path: {filepath}")
    
    # 4. Determinar tipo MIME
    media_type = "application/octet-stream"
    if safe_name.endswith(".pdf"):
        media_type = "application/pdf"
    elif safe_name.endswith(".xlsx"):
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif safe_name.endswith(".csv"):
        media_type = "text/csv"
    elif safe_name.endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    
    # 5. Ler arquivo e enviar com headers corretos
    with open(filepath, "rb") as f:
        file_bytes = f.read()
    
    # 6. Validar que o conteúdo XLSX começa com assinatura ZIP (PK\x03\x04)
    if safe_name.endswith(".xlsx") and file_bytes[:4] != b'PK\x03\x04':
        logger.error(f"[DOWNLOAD] Arquivo XLSX corrompido (não começa com PK): {safe_name}")
        raise HTTPException(status_code=500, detail="Arquivo Excel corrompido no servidor")
    
    # 7. Validar que o conteúdo PDF começa com %PDF
    if safe_name.endswith(".pdf") and file_bytes[:5] != b'%PDF-':
        logger.error(f"[DOWNLOAD] Arquivo PDF corrompido: {safe_name}")
        raise HTTPException(status_code=500, detail="Arquivo PDF corrompido no servidor")
    
    from fastapi.responses import FileResponse
    return FileResponse(
        path=filepath,
        filename=safe_name,
        media_type=media_type,
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "X-Content-Type-Options": "nosniff",
        }
    )

# ─── Chat ────────────────────────────────────────────────────────────

@router.post("/chat", summary="Conversar com a IA do sistema")
@_ai_limiter.limit("20/minute")
def ai_chat(
    request: ChatRequest,
    req: Request,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not GEMINI_API_KEY:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY não configurada no servidor.")
    
    # Propagar contexto do usuário para as ferramentas da IA (C3)
    if current_user.api_key:
        set_ai_user_context(current_user.api_key)
    
    # Buscar ou criar sessão
    session = None
    if request.session_id:
        session = (
            db.query(ChatSession)
            .filter(ChatSession.id == request.session_id, ChatSession.user_id == current_user.id)
            .first()
        )
    
    if not session:
        session = ChatSession(user_id=current_user.id, titulo="Nova conversa")
        db.add(session)
        db.commit()
        db.refresh(session)
    
    # Montar mensagem do usuário com contexto de arquivo se houver
    user_message = request.message
    if request.file_context:
        user_message = f"[O usuário enviou um documento. Conteúdo extraído abaixo]\n\n{request.file_context}\n\n[Mensagem do usuário]: {request.message}"
    
    # Salvar mensagem do usuário no banco
    user_msg = ChatMessage(session_id=session.id, role="user", content=request.message)
    db.add(user_msg)
    db.commit()
    
    # Injetar dinamicamente a arquitetura de rotas da aplicação  
    openapi_schema = req.app.openapi()
    endpoints_info = []
    for path, methods in openapi_schema.get("paths", {}).items():
        if path.startswith("/api/"):
            for method, info in methods.items():
                endpoints_info.append(f"[{method.upper()}] {path} - {info.get('summary', '')}")
                
    dynamic_system_prompt = SYSTEM_PROMPT + "\n\n### LISTA DE ENDPOINTS (APIs LOCAIS) QUE VOCÊ PODE CHAMAR VIA 'call_internal_api':\n" + "\n".join(endpoints_info)
    
    # Prepara o modelo com as ferramentas
    try:
        model = genai.GenerativeModel(
            model_name="gemini-2.5-flash",
            tools=AVAILABLE_TOOLS,
            system_instruction=dynamic_system_prompt
        )
        
        # Construir histórico da conversa do banco de dados
        history = []
        db_messages = (
            db.query(ChatMessage)
            .filter(ChatMessage.session_id == session.id)
            .order_by(ChatMessage.created_at)
            .all()
        )
        
        # Excluir a última mensagem (que acabamos de salvar) pois será enviada como mensagem atual
        if db_messages:
            db_messages = db_messages[:-1]
        
        for msg in db_messages:
            role = "user" if msg.role == "user" else "model"
            
            # Previne roles consecutivos iguais
            if history and history[-1]["role"] == role:
                history[-1]["parts"][0]["text"] += f"\n\n{msg.content}"
            else:
                history.append({"role": role, "parts": [{"text": msg.content}]})
        
        # Segurança: O Gemini NÃO aceita que o 'history' termine com 'user'
        last_message = user_message
        if history and history[-1]["role"] == "user":
            last_popped = history.pop()
            last_message = last_popped["parts"][0]["text"] + f"\n\n{last_message}"
        
        # Inicia o chat
        chat = model.start_chat(history=history)
        
        # Envia a mensagem principal
        response = chat.send_message(last_message)
        
        # Loop de function calling com tratamento robusto de erros
        max_iterations = 10
        iteration = 0
        while iteration < max_iterations:
            iteration += 1
            
            # Verificar se a resposta tem candidatos válidos
            if not response.candidates:
                break
            
            candidate = response.candidates[0]
            
            # Verificar finish_reason para erros do Gemini
            finish_reason = None
            try:
                finish_reason = candidate.finish_reason
            except:
                pass
            
            # MALFORMED_FUNCTION_CALL ou outros erros — retentar sem tools
            if finish_reason and str(finish_reason) not in ('0', '1', 'STOP', 'FinishReason.STOP', 'MAX_TOKENS', 'FinishReason.MAX_TOKENS'):
                # Tentar reenviar sem tools para obter resposta em texto
                try:
                    model_no_tools = genai.GenerativeModel(
                        model_name="gemini-2.5-flash",
                        system_instruction=dynamic_system_prompt
                    )
                    chat_retry = model_no_tools.start_chat(history=history)
                    response = chat_retry.send_message(last_message)
                except:
                    pass
                break
            
            # Verificar se há function_call
            function_call = None
            try:
                if candidate.content and candidate.content.parts:
                    for part in candidate.content.parts:
                        if part.function_call and part.function_call.name:
                            function_call = part.function_call
                            break
            except:
                break
            
            if not function_call:
                break
                
            func_name = function_call.name
            
            # Extract arguments
            args = {}
            if hasattr(function_call, 'args'):
                for key, val in function_call.args.items():
                    args[key] = val
            
            # Execute local function
            if func_name in TOOL_FUNCTIONS:
                func = TOOL_FUNCTIONS[func_name]
                try:
                    result = func(**args)
                except Exception as e:
                    result = json.dumps({"error": str(e)})
            else:
                result = json.dumps({"error": f"Função {func_name} não encontrada."})
            
            # Send result back to Gemini
            try:
                response_struct = Struct()
                response_struct.update({"result": result})
                
                function_response = genai_protos.Part(
                    function_response=genai_protos.FunctionResponse(
                        name=func_name,
                        response=response_struct
                    )
                )
                response = chat.send_message(function_response)
            except Exception as e:
                # Se falhar ao enviar resposta da função, quebrar o loop
                break
            
        # Obter texto final com segurança
        final_text = ""
        try:
            if response.candidates and response.candidates[0].content and response.candidates[0].content.parts:
                for part in response.candidates[0].content.parts:
                    if hasattr(part, 'text') and part.text:
                        final_text += part.text
        except:
            pass
        
        if not final_text:
            final_text = "Desculpe, não consegui processar sua solicitação no momento. Tente reformular a pergunta ou tente novamente."
        
        # Salvar resposta da IA no banco
        ai_msg = ChatMessage(session_id=session.id, role="model", content=final_text)
        db.add(ai_msg)
        
        # Atualizar título da sessão se for a primeira mensagem
        msg_count = db.query(ChatMessage).filter(ChatMessage.session_id == session.id).count()
        if msg_count <= 2:  # user + model = primeiras 2 mensagens
            # Usar as primeiras palavras da mensagem do usuário como título
            titulo = request.message[:60]
            if len(request.message) > 60:
                titulo += "..."
            session.titulo = titulo
        
        db.commit()
        
        return {
            "role": "model",
            "content": final_text,
            "session_id": session.id
        }
        
    except Exception as e:
        logging.exception("Erro interno da IA")
        raise HTTPException(status_code=500, detail="Erro interno da IA. Tente novamente.")
    finally:
        clear_ai_user_context()

